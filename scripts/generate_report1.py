import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_report1(output_path):
    doc = docx.Document()
    
    # Page Setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("AcadPrep — Project Introduction | Page ")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = RGBColor(128, 128, 128)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # --- COVER / TITLE ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_inst.add_run("FPT UNIVERSITY\nDEPARTMENT OF SOFTWARE ENGINEERING")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    p_inst.paragraph_format.space_after = Pt(36)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_main = p_title.add_run("PROJECT REPORT\n")
    r_main.font.size = Pt(18)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0, 51, 102)
    r_sub = p_title.add_run("Report 1 – Project Introduction\n\n")
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(70, 70, 70)
    
    r_proj = p_title.add_run("Project: AcadPrep — Online TOEIC Preparation Platform\n(Microservices Architecture)")
    r_proj.font.size = Pt(13)
    r_proj.font.bold = True
    r_proj.font.color.rgb = RGBColor(30, 100, 180)
    p_title.paragraph_format.space_after = Pt(60)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_d = p_date.add_run("– Semester 8 | October 2026 –")
    r_d.font.size = Pt(11)
    r_d.font.italic = True
    r_d.font.color.rgb = RGBColor(100, 100, 100)
    p_date.paragraph_format.space_after = Pt(40)

    doc.add_page_break()

    # --- TABLE OF CONTENTS (Static Representation) ---
    p_toc_head = doc.add_paragraph()
    r = p_toc_head.add_run("Table of Contents")
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    p_toc_head.paragraph_format.space_after = Pt(12)

    toc_items = [
        ("I. Record of Changes", "2"),
        ("II. Project Introduction", "3"),
        ("    1. Overview", "3"),
        ("        1.1 Project Information", "3"),
        ("        1.2 Project Team", "3"),
        ("    2. Product Background", "4"),
        ("    3. Existing Systems", "5"),
        ("        3.1 Study4 (study4.com)", "5"),
        ("        3.2 Prep.vn (prep.vn)", "6"),
        ("        3.3 Comparative Analysis & Key Takeaways", "6"),
        ("    4. Business Opportunity", "7"),
        ("    5. Software Product Vision", "8"),
        ("    6. Project Scope & Limitations", "8"),
        ("        6.1 Major Features (FE-01 to FE-08)", "8"),
        ("        6.2 Limitations & Exclusions", "10"),
    ]

    for item, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(item)
        r1.font.size = Pt(10.5)
        # Tab leader effect
        dots_count = max(5, 75 - len(item) * 2)
        r2 = p.add_run(" " + "." * dots_count + " ")
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(180, 180, 180)
        r3 = p.add_run(page)
        r3.font.size = Pt(10.5)
        r3.font.bold = True

    doc.add_page_break()

    # --- SECTION I: RECORD OF CHANGES ---
    h1 = doc.add_paragraph()
    r = h1.add_run("I. Record of Changes")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)

    t_changes = doc.add_table(rows=2, cols=4)
    t_changes.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_changes)
    
    headers = ['Date', 'A* M, D', 'In charge', 'Change Description']
    for idx, text in enumerate(headers):
        cell = t_changes.cell(0, idx)
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    # Initial row
    row_data = ['09/10/2026', 'A', 'AcadPrep Team', 'Initial drafting and completion of Report 1: Project Introduction, Team, Product Background, Existing Systems, Business Opportunity, Vision, Scope & Major Features.']
    for idx, text in enumerate(row_data):
        cell = t_changes.cell(1, idx)
        set_cell_background(cell, "F9FBFD")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 3 else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(9.5)

    p_note = doc.add_paragraph()
    r_n = p_note.add_run("*A - Added    M - Modified    D - Deleted")
    r_n.font.size = Pt(9)
    r_n.font.italic = True
    r_n.font.color.rgb = RGBColor(120, 120, 120)
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(16)

    # --- SECTION II: PROJECT INTRODUCTION ---
    h2 = doc.add_paragraph()
    r = h2.add_run("II. Project Introduction")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)

    # 1. Overview
    h3_1 = doc.add_paragraph()
    r = h3_1.add_run("1. Overview")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)
    h3_1.paragraph_format.space_before = Pt(10)
    h3_1.paragraph_format.space_after = Pt(6)

    # 1.1 Project Information
    h4_1 = doc.add_paragraph()
    r = h4_1.add_run("1.1 Project Information")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_1.paragraph_format.space_before = Pt(6)
    h4_1.paragraph_format.space_after = Pt(4)

    proj_infos = [
        ("Project Name:", "AcadPrep — Online TOEIC Preparation Platform"),
        ("Project Code:", "ACADPREP"),
        ("Group Name:", "PRN232_Group04 (Semester 8 - PRN232)"),
        ("Software Type:", "Web Application (Distributed Microservices Architecture: ASP.NET Core 9, Razor Pages, YARP API Gateway, EF Core, RabbitMQ, Redis, VNPay Payment Gateway)"),
    ]
    for label, val in proj_infos:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"• {label} ")
        r1.font.bold = True
        r2 = p.add_run(val)

    # 1.2 Project Team
    h4_2 = doc.add_paragraph()
    r = h4_2.add_run("1.2 Project Team")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_2.paragraph_format.space_before = Pt(10)
    h4_2.paragraph_format.space_after = Pt(6)

    t_team = doc.add_table(rows=7, cols=4)
    t_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_team)
    
    team_headers = ['Full Name', 'Role & Service Responsibility', 'Email', 'Student ID / Note']
    for idx, text in enumerate(team_headers):
        cell = t_team.cell(0, idx)
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    team_members = [
        ('Instructor / Lecturer', 'Project Supervisor', 'lecturer@fe.edu.vn', 'FPT University'),
        ('Đặng Thanh Liêm', 'Team Leader / Exam Engine Service Lead (Core Test Engine, Timed Scoring, Attempts History)', 'liemdt.ce190697@gmail.com', 'CE190697 (Leader)'),
        ('Dương Châu Đăng Khoa', 'Member / Identity Service Lead (Authentication, OAuth2, JWT Security, RBAC)', 'khoadcd@fe.edu.vn', 'Member'),
        ('Hồng Tuấn Nguyên', 'Member / Content CMS Lead (Exam Bank, Questions, Passages, Content Publishing)', 'nguyenht@fe.edu.vn', 'Member'),
        ('Hứa Khánh Duy', 'Member / Billing & Notification Lead (Freemium Quotas, VNPay Checkout, IPN, Notifications)', 'duyhk@fe.edu.vn', 'Member'),
        ('Đỗ Hồ Gia Huy', 'Member / Learning & Analytics Lead (Spaced Repetition Flashcards, Streak, Leaderboard, Analytics)', 'huyddh@fe.edu.vn', 'Member'),
    ]

    for row_idx, row_data in enumerate(team_members, start=1):
        bg_col = "FFFFFF" if row_idx % 2 != 0 else "F4F8FC"
        for col_idx, text in enumerate(row_data):
            cell = t_team.cell(row_idx, col_idx)
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            if row_idx == 2 and col_idx == 0:
                run.font.bold = True
            run.font.size = Pt(9.5)

    # 2. Product Background
    h3_2 = doc.add_paragraph()
    r = h3_2.add_run("2. Product Background")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)
    h3_2.paragraph_format.space_before = Pt(14)
    h3_2.paragraph_format.space_after = Pt(6)

    bg_texts = [
        "In the contemporary global academic and corporate landscape, the Test of English for International Communication (TOEIC) serves as an indispensable benchmark for assessing workplace English proficiency. At Vietnamese universities (including FPT University, VNU, and polytechnics) and multinational enterprises, achieving a TOEIC score of 500 to 850+ is a mandatory graduation exit criterion and a pivotal prerequisite for career recruitment and advancement.",
        "Despite the vital importance of TOEIC certification, the overwhelming majority of learners encounter substantial hurdles during their preparation journey:",
        "• High Costs and Inflexible Schedules: Traditional brick-and-mortar English preparation centers impose high tuition fees and rigid schedules, which are ill-suited for busy university students and working professionals.",
        "• Passive Learning from Physical Textbooks: Self-study using conventional printed books lacks interactive audio synchronization, instantaneous scoring, and detailed explanation breakdowns. Learners often spend excessive time manually cross-checking answer keys without understanding why an option was incorrect.",
        "• Inefficient Vocabulary Memorization: Most learners struggle to retain technical business vocabulary due to passive rote memorization without scientifically proven methods like Spaced Repetition (SRS).",
        "• Monolithic Legacy Test Websites: Many available test-prep websites suffer from sluggish response times during peak hours, uninspiring user interfaces, lack of real-time progress analytics, and absence of seamless micro-transaction subscription plans.",
        "To decisively overcome these challenges, the AcadPrep team proposed and developed an all-in-one, highly scalable, and user-centric online TOEIC learning platform. Built upon a distributed Microservices Architecture, AcadPrep delivers an authentic 2-hour full test simulation, targeted part-by-part drills, intelligent SRS-based flashcard memorization, real-time performance analytics, and a seamless Freemium monetization model integrated with VNPay. This empowers learners to study anytime, anywhere, with maximum efficiency and at an affordable cost."
    ]
    for b_txt in bg_texts:
        p = doc.add_paragraph(b_txt)
        if b_txt.startswith("•"):
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(3)
        else:
            p.paragraph_format.space_after = Pt(6)

    # 3. Existing Systems
    h3_3 = doc.add_paragraph()
    r = h3_3.add_run("3. Existing Systems")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)
    h3_3.paragraph_format.space_before = Pt(14)
    h3_3.paragraph_format.space_after = Pt(6)

    p_exist_intro = doc.add_paragraph("To design an effective architecture and competitive feature set, our team conducted a thorough evaluation of prominent existing test preparation platforms in the market:")
    p_exist_intro.paragraph_format.space_after = Pt(6)

    # 3.1 Study4
    h4_s1 = doc.add_paragraph()
    r = h4_s1.add_run("3.1 Study4 (https://study4.com)")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_s1.paragraph_format.space_before = Pt(8)
    h4_s1.paragraph_format.space_after = Pt(4)

    s1_texts = [
        "• System Overview: Study4 is a well-known Vietnamese EdTech platform offering online mock tests for TOEIC, IELTS, and JLPT, alongside spaced-repetition flashcard tools.",
        "• Key System Actors: Guest (trial users), Learner (registered users/subscribers), Content Manager, Administrator.",
        "• Major Features: Full-length TOEIC test rooms, part-by-part mini tests, community answer discussion forums, flashcard deck creation, and audio transcript display.",
        "• Strengths (Pros): Clean web interface, extensive repository of past ETS exam papers, active learner discussion section below each question.",
        "• Weaknesses (Cons): The system is built predominantly as a monolithic web application, which experiences noticeable latency and occasional service degradation during evening peak hours. Premium course packages are expensive, and analytical breakdowns between sub-skills (e.g., Inference vs. Detail questions in Part 7) are relatively limited."
    ]
    for t in s1_texts:
        p = doc.add_paragraph(t)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)

    # 3.2 Prep.vn
    h4_s2 = doc.add_paragraph()
    r = h4_s2.add_run("3.2 Prep.vn (https://prep.vn)")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_s2.paragraph_format.space_before = Pt(8)
    h4_s2.paragraph_format.space_after = Pt(4)

    s2_texts = [
        "• System Overview: Prep.vn is a comprehensive smart learning platform focused on standardized language exams with interactive video courses, mock test rooms, and AI-assisted grading.",
        "• Key System Actors: Learner, Academic Mentor/Teacher, Consultant, System Administrator.",
        "• Major Features: Structured video learning roadmaps, virtual testing rooms with anti-cheat timers, AI Virtual Speaking/Writing grading room, and scheduled study reminders.",
        "• Strengths (Pros): High pedagogical quality, well-curated study pathways, excellent marketing, and modern branding.",
        "• Weaknesses (Cons): High barrier to entry with costly course bundles; lacks a dedicated lightweight freemium drill-oriented workflow for rapid daily exam practice and Spaced Repetition vocabulary lookup directly from mock tests."
    ]
    for t in s2_texts:
        p = doc.add_paragraph(t)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)

    # 3.3 Comparative Analysis & Key Takeaways
    h4_s3 = doc.add_paragraph()
    r = h4_s3.add_run("3.3 Comparative Analysis & Key Takeaways for AcadPrep")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_s3.paragraph_format.space_before = Pt(8)
    h4_s3.paragraph_format.space_after = Pt(6)

    t_comp = doc.add_table(rows=6, cols=4)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_comp)

    comp_headers = ['Evaluation Criteria', 'Study4', 'Prep.vn', 'AcadPrep (Our System)']
    for idx, text in enumerate(comp_headers):
        cell = t_comp.cell(0, idx)
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)

    comp_data = [
        ('Core Architecture', 'Monolith Web Application', 'Service-Oriented / Cloud Monolith', 'Independent Microservices (Clean Arch, YARP, RabbitMQ, Redis)'),
        ('Freemium Model', 'Limited Free Tests (Ad-supported)', 'Minimal Free Content (Course-gated)', 'Flexible Freemium (Daily free practice + affordable VNPay Pro tier)'),
        ('Vocabulary & SRS', 'Standard Flashcards', 'Vocabulary Lists per Video', 'In-Test Contextual Vocabulary Notebook + SuperMemo-based SRS Rating'),
        ('Test Engine & UX', 'Full Test + Mini Test', 'Full Test Room', 'Full Test 7 Parts + Part Drill + Resume + Granular Detailed Explanations'),
        ('Performance & Scale', 'Moderate (Peak latency spikes)', 'High (Heavy frontend bandwidth)', 'High Resilience & Zero-downtime microservice scaling'),
    ]

    for r_idx, r_data in enumerate(comp_data, start=1):
        bg_col = "FFFFFF" if r_idx % 2 != 0 else "F4F8FC"
        for c_idx, text in enumerate(r_data):
            cell = t_comp.cell(r_idx, c_idx)
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            if c_idx == 3:
                run.font.bold = True
            run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Business Opportunity
    h3_4 = doc.add_paragraph()
    r = h3_4.add_run("4. Business Opportunity")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)
    h3_4.paragraph_format.space_before = Pt(14)
    h3_4.paragraph_format.space_after = Pt(6)

    biz_texts = [
        "The market demand for standardized English assessment in Vietnam has surged exponentially over the past five years. Annually, over 500,000 university students and early-career job seekers register for TOEIC examinations across major cities (Hanoi, Ho Chi Minh City, Da Nang, Can Tho). Despite high demand, there is a clear market gap between expensive physical tutoring centers (costing 3,000,000–8,000,000 VND per course) and low-engagement static PDF mock tests.",
        "AcadPrep seizes this lucrative opportunity by introducing an agile, cost-efficient, and highly scalable online preparation platform characterized by:",
        "1. High-Conversion Freemium Model: Unregistered and free-tier users can experience core features (taking selected practice parts, saving initial vocabulary, and testing platform ergonomics). When users desire unlimited full-length mock exams, comprehensive question rationales, and advanced score projections, they can instantly upgrade to Pro via VNPay with affordable micro-packages (e.g., 49,000 VND/month or 129,000 VND/quarter).",
        "2. Cost-Effective Microservices Infrastructure: By leveraging lightweight .NET 9 services, Redis distributed caching, and asynchronous RabbitMQ event processing, AcadPrep maintains low server overhead while easily supporting concurrent spikes during exam seasons.",
        "3. High User Retention via Gamification & SRS: Study streaks, achievement badges, and spaced repetition review schedules cultivate long-term daily learning habits, generating recurring user engagement and word-of-mouth referral growth.",
        "4. Seamless Domestic Payment Integration: Direct integration with VNPay (supporting VNPay QR, domestic ATM cards, and mobile banking apps) removes payment friction for Vietnamese students who often do not possess international credit cards."
    ]
    for b in biz_texts:
        p = doc.add_paragraph(b)
        p.paragraph_format.space_after = Pt(6)

    # 5. Software Product Vision
    h3_5 = doc.add_paragraph()
    r = h3_5.add_run("5. Software Product Vision")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)
    h3_5.paragraph_format.space_before = Pt(14)
    h3_5.paragraph_format.space_after = Pt(6)

    p_v_box = doc.add_paragraph()
    p_v_box.paragraph_format.left_indent = Inches(0.2)
    p_v_box.paragraph_format.right_indent = Inches(0.2)
    p_v_box.paragraph_format.space_before = Pt(4)
    p_v_box.paragraph_format.space_after = Pt(10)
    
    r_v1 = p_v_box.add_run("“For learners and working professionals ")
    r_v1.font.italic = True
    r_v1.font.size = Pt(11)
    
    r_v2 = p_v_box.add_run("who want an authentic, self-paced, and data-driven approach to mastering the TOEIC examination, ")
    r_v2.font.italic = True
    
    r_v3 = p_v_box.add_run("AcadPrep is a modern, high-performance web platform ")
    r_v3.font.bold = True
    r_v3.font.color.rgb = RGBColor(0, 51, 102)
    
    r_v4 = p_v_box.add_run("built on a distributed microservices architecture that provides real-time exam simulations (Full Test 7 Parts & Part Drills), contextual Spaced Repetition (SRS) vocabulary flashcards, in-depth analytical performance tracking, and seamless domestic VNPay payment for Pro access. ")
    r_v4.font.italic = True
    
    r_v5 = p_v_box.add_run("Unlike traditional monolithic websites or expensive offline test centers, ")
    r_v5.font.italic = True
    
    r_v6 = p_v_box.add_run("AcadPrep delivers zero-friction accessibility, uninterrupted testing with automatic countdown and resume, instant explanations, and habit-forming gamification that accelerates score gains and ensures examination success.”")
    r_v6.font.bold = True

    # 6. Project Scope & Limitations
    h3_6 = doc.add_paragraph()
    r = h3_6.add_run("6. Project Scope & Limitations")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 70, 140)
    h3_6.paragraph_format.space_before = Pt(14)
    h3_6.paragraph_format.space_after = Pt(6)

    p_scope_intro = doc.add_paragraph(
        "To ensure precise project execution within the academic semester and align with the distributed microservices architecture, the project scope, key functional capabilities, and explicit boundaries have been strictly established."
    )
    p_scope_intro.paragraph_format.space_after = Pt(6)

    # 6.1 Major Features
    h4_6_1 = doc.add_paragraph()
    r = h4_6_1.add_run("6.1 Major Features")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_6_1.paragraph_format.space_before = Pt(8)
    h4_6_1.paragraph_format.space_after = Pt(6)

    features = [
        ("FE-01: User Identity & Access Management",
         "Provides secure user onboarding and authentication. Includes Email/Password registration with 6-digit OTP email verification, Google OAuth2 single sign-on, JWT-based stateless session management, secure password reset, profile management, and Role-Based Access Control (Admin, Moderator, Learner)."),
        ("FE-02: TOEIC Exam Engine & Simulation",
         "Supports two testing modes: Full Test (200 questions across 7 Parts, 120-minute countdown, audio player with play restrictions) and Practice by Part (custom part selection without strict time pressure). Features real-time state persistence and session resumption in case of unexpected network disconnection."),
        ("FE-03: Instant Scoring & Comprehensive Answer Review",
         "Automatically grades submitted tests according to standardized TOEIC score conversion tables, calculates scaled Listening and Reading sub-scores, presents comprehensive question-by-question review with answer rationales, and catalogs incorrect answers into a dedicated review repository."),
        ("FE-04: Smart Vocabulary Notebook & Spaced Repetition (SRS)",
         "Enables learners to look up words in context and save them directly from reading passages into personal vocabulary notebooks. Integrates a SuperMemo-based Spaced Repetition algorithm (SRS Flashcards) with recall difficulty ratings (1 to 5) to optimize long-term memory retention."),
        ("FE-05: Learning Analytics, Gamification & Leaderboard",
         "Provides visual performance dashboards illustrating Listening vs. Reading strengths/weaknesses, score progress trends over time, daily study streak tracking (with UTC+7 timezone auto-reset), milestone achievement badges, and competitive learner leaderboards."),
        ("FE-06: Back-Office Content Management System (CMS)",
         "Empowers Content Moderators and Administrators to author, update, structure, and publish TOEIC exams, question banks (with audio files, transcripts, images, and explanations), and multi-question reading passages with referential integrity validation."),
        ("FE-07: Subscription Billing & VNPay Integration",
         "Implements a Freemium model with quota enforcement (daily limits on full tests, practice sessions, and vocabulary saves for free tier). Provides VNPay checkout integration (HMAC-SHA512 checksum, IPN webhook handling, idempotent transaction processing) for instant Pro plan activation."),
        ("FE-08: Real-Time System Notifications & Administrative Reporting",
         "Delivers in-app notifications for exam completion, daily study streak milestones, and subscription status. Equips administrators with comprehensive reporting dashboards covering new registrations, retention, revenue, and exam completion metrics.")
    ]

    for f_code, f_desc in features:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        r_fc = p.add_run(f"• {f_code}:\n")
        r_fc.font.bold = True
        r_fc.font.color.rgb = RGBColor(0, 51, 102)
        r_fd = p.add_run(f"  {f_desc}")

    # 6.2 Limitations & Exclusions
    h4_6_2 = doc.add_paragraph()
    r = h4_6_2.add_run("6.2 Limitations & Exclusions")
    r.font.size = Pt(11.5)
    r.font.bold = True
    h4_6_2.paragraph_format.space_before = Pt(10)
    h4_6_2.paragraph_format.space_after = Pt(6)

    limitations = [
        ("LI-01: Focus on TOEIC Listening & Reading Only",
         "The current system exclusively supports the standard 2-skill TOEIC Listening & Reading format (Parts 1 to 7). TOEIC Speaking & Writing evaluations requiring subjective human grading or certified examiners are excluded from this release."),
        ("LI-02: Payment Gateway Boundary (VNPay in VND)",
         "Payment processing is strictly integrated with the VNPay gateway supporting Vietnamese Dong (VND) transactions (VNPay-QR, local bank ATM/iBanking). International payment gateways (Stripe, PayPal) and multi-currency billing are not included in Phase 1."),
        ("LI-03: Web Application Delivery (No Native Mobile Apps)",
         "AcadPrep is delivered as a fully responsive Web Application optimized for desktop, tablet, and mobile browsers. Native mobile binaries (.apk for Android or .ipa for iOS) are outside the project boundary."),
        ("LI-04: No Human 1-on-1 Tutoring / Live Video Streaming",
         "AcadPrep is designed as an automated, self-paced learning and mock testing system; real-time video conferencing, live classroom broadcasting, and 1-on-1 human tutoring sessions are excluded from the system scope."),
        ("LI-05: Single Organization / Non-Multi-Tenant Deployment",
         "The platform is deployed as a single-tenant e-learning system for AcadPrep users; white-labeling and multi-tenant school partitioning are reserved for future commercial scaling.")
    ]

    for l_code, l_desc in limitations:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        r_lc = p.add_run(f"• {l_code}:\n")
        r_lc.font.bold = True
        r_lc.font.color.rgb = RGBColor(150, 40, 40)
        r_ld = p.add_run(f"  {l_desc}")

    # Save document
    doc.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == '__main__':
    build_report1('docs/Report1_Project Introduction.docx')
